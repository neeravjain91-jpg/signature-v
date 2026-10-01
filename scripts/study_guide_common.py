#!/usr/bin/env python3
"""
SIGNATURE VMAKE — Study Guide Common Formatting & Helper Library
Provides unified styling, callout boxes, table generators, and figure insertion for python-docx.
"""

import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Color Constants
C_NAVY_DARK = RGBColor(11, 17, 32)      # #0B1120
C_BLUE_PRIMARY = RGBColor(30, 64, 175)  # #1E40AF
C_CYAN_ACCENT = RGBColor(2, 132, 199)   # #0284C7
C_EMERALD = RGBColor(16, 185, 129)      # #10B981
C_AMBER = RGBColor(217, 119, 6)         # #D97706
C_ROSE = RGBColor(225, 29, 72)          # #E11D48
C_PURPLE = RGBColor(126, 34, 206)       # #7E22CE
C_TEXT_DARK = RGBColor(30, 41, 59)      # #1E293B
C_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B

HEX_BG_LIGHT_BLUE = "F0F9FF"
HEX_BG_LIGHT_AMBER = "FFFBEB"
HEX_BG_LIGHT_GREEN = "F0FDF4"
HEX_BG_LIGHT_PURPLE = "FAF5FF"
HEX_BG_LIGHT_GRAY = "F8FAFC"


def setup_document_styles(doc):
    """Configures page geometry, fonts, margins, headers and footers."""
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.page_width = Inches(8.27)   # A4 Width
        section.page_height = Inches(11.69) # A4 Height

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SIGNATURE VMAKE: AI-Powered Verification — Complete Study Guide | B.Tech 3rd Year SRS")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = C_TEXT_MUTED

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential — Academic Examination & Project Viva Defense Manual  •  IEEE Std 830-1998")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8)
        frun.font.color.rgb = C_TEXT_MUTED


def add_part_heading(doc, part_num: int, part_title: str):
    """Adds a major Part Divider on a fresh page."""
    if part_num > 1:
        doc.add_page_break()

    p_pre = doc.add_paragraph()
    p_pre.space_before = Pt(28)
    p_pre.space_after = Pt(2)
    r_pre = p_pre.add_run(f"PART {part_num}".upper())
    r_pre.font.name = "Calibri"
    r_pre.font.size = Pt(12)
    r_pre.font.bold = True
    r_pre.font.color.rgb = C_CYAN_ACCENT

    p = doc.add_heading(part_title, level=1)
    p.space_before = Pt(0)
    p.space_after = Pt(14)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = C_BLUE_PRIMARY


def add_chapter_heading(doc, chapter_num: int, chapter_title: str):
    """Adds a Chapter Heading."""
    p = doc.add_heading(f"Chapter {chapter_num} — {chapter_title}", level=2)
    p.space_before = Pt(18)
    p.space_after = Pt(8)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(14.5)
        run.font.bold = True
        run.font.color.rgb = C_NAVY_DARK


def add_section_heading(doc, section_title: str):
    """Adds a Section sub-heading."""
    p = doc.add_heading(section_title, level=3)
    p.space_before = Pt(12)
    p.space_after = Pt(4)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = C_CYAN_ACCENT


def add_body_p(doc, text: str, bold_prefix: str = None):
    """Adds standard body paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = C_TEXT_DARK
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = C_TEXT_DARK
    return p


def add_bullet_p(doc, text: str, bold_prefix: str = None):
    """Adds a bullet point paragraph."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = C_TEXT_DARK
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = C_TEXT_DARK
    return p


def add_callout(doc, box_type: str, title: str, content: str):
    """
    Renders an academic callout container box.
    box_type: 'VIVA TIP', 'COMMON MISTAKE', 'REMEMBER', 'TECHNICAL DEEP DIVE'
    """
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    border_colors = {
        "VIVA TIP": ("0284C7", HEX_BG_LIGHT_BLUE, C_CYAN_ACCENT),
        "COMMON MISTAKE": ("E11D48", HEX_BG_LIGHT_AMBER, C_ROSE),
        "REMEMBER": ("10B981", HEX_BG_LIGHT_GREEN, C_EMERALD),
        "TECHNICAL DEEP DIVE": ("7E22CE", HEX_BG_LIGHT_PURPLE, C_PURPLE)
    }
    hex_border, hex_bg, title_col = border_colors.get(box_type, ("0284C7", HEX_BG_LIGHT_BLUE, C_CYAN_ACCENT))

    cell = tbl.cell(0, 0)
    cell.width = Inches(6.47)
    tcPr = cell._tc.get_or_add_tcPr()

    # Background shading
    shd_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_bg}"/>'
    tcPr.append(parse_xml(shd_xml))

    # Left thick colored border, no other borders
    borders_xml = f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="28" w:space="0" w:color="{hex_border}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    r_tag = p.add_run(f"[{box_type.upper()}]  {title}\n")
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(10.5)
    r_tag.font.bold = True
    r_tag.font.color.rgb = title_col

    r_body = p.add_run(content)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = C_TEXT_DARK

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(2)


def add_dual_level_explanation(doc, beginner_text: str, technical_text: str):
    """
    Renders structured dual-level teaching explanation:
    LEVEL 1 - Beginner Intuition & Real-World Analogy
    LEVEL 2 - Engineering & Mathematical Precision
    """
    # Level 1 Card
    tbl1 = doc.add_table(rows=1, cols=1)
    tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1 = tbl1.cell(0, 0)
    c1.width = Inches(6.47)
    tcPr1 = c1._tc.get_or_add_tcPr()
    tcPr1.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_BG_LIGHT_GREEN}"/>'))
    tcPr1.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="20" w:space="0" w:color="10B981"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>'))
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.line_spacing = 1.15
    r1_h = p1.add_run("LEVEL 1 — Intuitive Beginner Explanation (The 'Explain Like I'm Five' Analogy):\n")
    r1_h.font.name = "Calibri"
    r1_h.font.size = Pt(10)
    r1_h.font.bold = True
    r1_h.font.color.rgb = C_EMERALD
    r1_b = p1.add_run(beginner_text)
    r1_b.font.name = "Calibri"
    r1_b.font.size = Pt(9.5)
    r1_b.font.color.rgb = C_TEXT_DARK

    # Level 2 Card
    tbl2 = doc.add_table(rows=1, cols=1)
    tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    c2 = tbl2.cell(0, 0)
    c2.width = Inches(6.47)
    tcPr2 = c2._tc.get_or_add_tcPr()
    tcPr2.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_BG_LIGHT_BLUE}"/>'))
    tcPr2.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="20" w:space="0" w:color="1E40AF"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>'))
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_before = Pt(4)
    p2.paragraph_format.space_after = Pt(4)
    p2.paragraph_format.line_spacing = 1.15
    r2_h = p2.add_run("LEVEL 2 — Engineering & Mathematical Deep-Dive (Technical Rigor for Examiners):\n")
    r2_h.font.name = "Calibri"
    r2_h.font.size = Pt(10)
    r2_h.font.bold = True
    r2_h.font.color.rgb = C_BLUE_PRIMARY
    r2_b = p2.add_run(technical_text)
    r2_b.font.name = "Calibri"
    r2_b.font.size = Pt(9.5)
    r2_b.font.color.rgb = C_TEXT_DARK

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(4)


def add_component_profile(doc, what: str, why: str, how: str, inputs: str, outputs: str, connections: str, location: str):
    """
    Renders the formal 7-point engineering component specification:
    WHAT, WHY, HOW, INPUT, OUTPUT, CONNECTIONS, REPOSITORY LOCATION.
    """
    tbl = doc.add_table(rows=7, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Inches(1.8), Inches(4.67)]
    rows = [
        ("WHAT IS IT?", what),
        ("WHY IS IT USED?", why),
        ("HOW DOES IT WORK?", how),
        ("INPUT RECEIVED", inputs),
        ("OUTPUT PRODUCED", outputs),
        ("SYSTEM CONNECTIONS", connections),
        ("REPOSITORY LOCATION", location)
    ]
    for row_idx, (label, val) in enumerate(rows):
        c0 = tbl.cell(row_idx, 0)
        c1 = tbl.cell(row_idx, 1)
        c0.width = col_widths[0]
        c1.width = col_widths[1]

        # Shading
        c0_bg = "1E293B" if row_idx == 0 else "F1F5F9"
        c1_bg = "0F172A" if row_idx == 0 else ("FFFFFF" if row_idx % 2 == 1 else "F8FAFC")
        c0._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{c0_bg}"/>'))
        c1._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{c1_bg}"/>'))

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(255, 255, 255) if row_idx == 0 else C_BLUE_PRIMARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor(255, 255, 255) if row_idx == 0 else C_TEXT_DARK

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(4)


def add_styled_table(doc, headers, data, col_widths=None):
    """Adds a standard data table with dark header row and alternating cell shading."""
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Header Row
    for col_idx, h in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        if col_widths and col_idx < len(col_widths):
            cell.width = col_widths[col_idx]
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>'))
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, rdata in enumerate(data, start=1):
        bg_hex = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(rdata):
            cell = tbl.cell(row_idx, col_idx)
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>'))
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(val))
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = C_BLUE_PRIMARY
            else:
                r.font.color.rgb = C_TEXT_DARK

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(4)


def add_figure(doc, img_path: Path, caption: str, width_inches: float = 6.2):
    """Embeds an image with an academic figure caption."""
    if not img_path.exists():
        p_err = doc.add_paragraph(f"[Image Missing: {img_path.name}]")
        p_err.runs[0].font.color.rgb = C_ROSE
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run().add_picture(str(img_path), width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(1)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Calibri"
    r_cap.font.size = Pt(8.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = C_TEXT_MUTED


ASSET_DIR = Path(os.environ.get('TEMP', 'C:/Users/ASUS/AppData/Local/Temp')) / 'vmake_ppt_assets'


def add_title_page(doc):
    """Generates a formal university B.Tech 3rd Year SRS-style title page."""
    # Top spacing
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(36)
    p_top.paragraph_format.space_after = Pt(12)
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_top.add_run("DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\nB.TECH 3RD YEAR PROJECT MANUAL & VIVA DEFENSE GUIDE")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(12)
    r_inst.font.bold = True
    r_inst.font.color.rgb = C_CYAN_ACCENT

    # Major Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("SIGNATURE VMAKE")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(32)
    r_title.font.bold = True
    r_title.font.color.rgb = C_BLUE_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(20)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("AI-Powered Offline Signature Verification & Banking Document Authentication System\nComprehensive Academic SRS, Architecture Manual & Viva Examiner Study Guide")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_NAVY_DARK

    # Academic Meta Box
    tbl = doc.add_table(rows=6, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_rows = [
        ("Specification Standard", "IEEE Std 830-1998 (Recommended Practice for Software Requirements Specifications)"),
        ("Approved Tech Stack", "Python 3.11  •  scikit-learn  •  Hugging Face Transformers  •  FastAPI  •  PostgreSQL"),
        ("Computer Vision Core", "OpenCV (Otsu Binarization, Zhang-Suen Thinning, 16 Structural & Geometric Descriptors)"),
        ("Machine Learning Models", "Track A: Random Forest (ROC-AUC 0.9424), SVM, Logistic  •  Track B: DeiT-Tiny (Hf)"),
        ("Evaluation Dataset", "CEDAR Offline Signature Benchmark (55 Writers, 2,640 Signatures, Writer-Disjoint)"),
        ("Testing & Quality Gate", "53 / 53 Pytest Suite Passing (100%)  •  16 / 16 Diagnostic Health Checks Verified")
    ]
    for idx, (label, val) in enumerate(meta_rows):
        c0 = tbl.cell(idx, 0)
        c1 = tbl.cell(idx, 1)
        c0.width = Inches(2.0)
        c1.width = Inches(4.47)
        c0._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>'))
        c1._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>'))
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = C_BLUE_PRIMARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = C_TEXT_DARK

    # Notice callout
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(28)
    p_note.paragraph_format.space_after = Pt(12)
    p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_n = p_note.add_run("CRITICAL ACADEMIC SEPARATION NOTICE:\nThis project is strictly distinct from the legacy Siamese metric-learning project SYNAPSE. The production champion is scikit-learn Random Forest operating on handcrafted multi-domain feature vectors, supplemented by a Hugging Face Vision Transformer (facebook/deit-tiny-patch16-224). All evaluations strictly enforce zero-leakage writer-disjoint partitioning on the CEDAR benchmark.")
    r_n.font.name = "Calibri"
    r_n.font.size = Pt(9)
    r_n.font.italic = True
    r_n.font.color.rgb = C_TEXT_MUTED

    doc.add_page_break()


def add_code_block(doc, code_str: str, caption: str = None):
    """Renders a monospaced syntax box with border and shading."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.47)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>'))
    tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="16" w:space="0" w:color="94A3B8"/><w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/></w:tcBorders>'))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05

    if caption:
        rc = p.add_run(f"// {caption}\n")
        rc.font.name = "Consolas"
        rc.font.size = Pt(8.5)
        rc.font.bold = True
        rc.font.color.rgb = C_BLUE_PRIMARY

    r = p.add_run(code_str.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(15, 23, 42)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(4)


def add_math_formula(doc, formula_str: str, explanation: str = None):
    """Renders a centered mathematical formula box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.47)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="EFF6FF"/>'))
    tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="3B82F6"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>'))

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)

    r_f = p.add_run(f"{formula_str}\n")
    r_f.font.name = "Cambria Math"
    r_f.font.size = Pt(11)
    r_f.font.bold = True
    r_f.font.color.rgb = C_BLUE_PRIMARY

    if explanation:
        r_exp = p.add_run(explanation)
        r_exp.font.name = "Calibri"
        r_exp.font.size = Pt(8.5)
        r_exp.font.italic = True
        r_exp.font.color.rgb = C_TEXT_MUTED

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(4)


print("study_guide_common library loaded successfully with title page and math helpers.")
